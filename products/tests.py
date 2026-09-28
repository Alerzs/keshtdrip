from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from products.models import Brand, Category, Product, ProductVariant
from products.price_import import PriceImportError, import_price_workbook


class ProductVariantTests(TestCase):
    def setUp(self):
        category = Category.objects.create(name="آبیاری تست", slug="variant-irrigation", tagline="آب")
        brand = Brand.objects.create(name="برند تست", slug="variant-brand")
        self.product = Product.objects.create(
            category=category,
            brand=brand,
            name="لوله قطره‌ای",
            slug="drip-pipe",
            blurb="لوله",
            description="توضیح",
            price=1000,
            unit="متر",
        )
        self.low = ProductVariant.objects.create(
            product=self.product,
            slug="drip-pipe-16",
            working_pressure="2 بار",
            diameter="16 میلی‌متر",
            price=100000,
        )
        self.high = ProductVariant.objects.create(
            product=self.product,
            slug="drip-pipe-20",
            working_pressure="4 بار",
            diameter="20 میلی‌متر",
            price=180000,
        )
        self.plain = Product.objects.create(
            category=category,
            brand=brand,
            name="کوزه",
            slug="olla",
            blurb="کوزه",
            description="توضیح",
            price=640000,
            unit="جفت",
        )

    def test_catalog_shows_each_pressure_and_diameter_as_its_own_box(self):
        response = self.client.get(reverse("catalog"))
        self.assertContains(response, "لوله قطره‌ای", count=3)
        self.assertContains(response, "۱۰۰٬۰۰۰")
        self.assertContains(response, "۱۸۰٬۰۰۰")
        self.assertContains(response, "کوزه", count=2)

    def test_product_page_prices_the_selected_combination(self):
        response = self.client.get(self.high.get_absolute_url())
        self.assertContains(response, "۱۸۰٬۰۰۰")
        self.assertContains(response, 'name="working_pressure"')
        self.assertContains(response, 'name="diameter"')
        self.assertContains(response, 'name="variant_slug"')
        self.assertContains(response, 'value="drip-pipe-20"')
        self.assertContains(response, "4 بار")
        self.assertContains(response, "20 میلی‌متر")

    def test_cart_uses_the_chosen_combination_price(self):
        self.client.post(
            reverse("add_to_cart", args=[self.product.slug]),
            {"quantity": 2, "variant_slug": "drip-pipe-20"},
        )
        self.client.post(
            reverse("add_to_cart", args=[self.product.slug]),
            {"quantity": 1, "variant_slug": "drip-pipe-16"},
        )
        response = self.client.get(reverse("cart"))
        self.assertContains(response, "قطر ۲۰ میلی‌متر، فشار کاری ۴ بار")
        self.assertContains(response, "قطر ۱۶ میلی‌متر، فشار کاری ۲ بار")
        self.assertContains(response, "drip-pipe-20")
        self.assertContains(response, "drip-pipe-16")
        self.assertContains(response, "۳۶۰٬۰۰۰")
        self.assertContains(response, "۴۶۰٬۰۰۰")

    def test_unknown_combination_is_rejected(self):
        response = self.client.post(
            reverse("add_to_cart", args=[self.product.slug]),
            {"quantity": 1, "variant_slug": "missing-variant"},
        )
        self.assertRedirects(response, self.product.get_absolute_url())
        cart = self.client.get(reverse("cart"))
        self.assertContains(cart, "سبد خرید")
        self.assertNotContains(cart, "لوله قطره‌ای")

    def test_cart_requires_the_variant_slug(self):
        response = self.client.post(
            reverse("add_to_cart", args=[self.product.slug]),
            {"quantity": 1, "working_pressure": "4 بار", "diameter": "20 میلی‌متر"},
        )
        self.assertRedirects(response, self.product.get_absolute_url())
        cart = self.client.get(reverse("cart"))
        self.assertNotContains(cart, "لوله قطره‌ای")


class SeoTests(TestCase):
    def test_home_targets_tape_and_irrigation_keywords(self):
        response = self.client.get(reverse("home"))
        self.assertContains(response, "خرید نوار تیپ و لوازم آبیاری | کشت‌دریپ")
        self.assertContains(response, "خرید نوار تیپ پلاکدار و درزدار")
        self.assertContains(response, 'rel="canonical"')
        self.assertContains(response, "application/ld+json")
        self.assertContains(response, "OnlineStore")
        html = response.content.decode()
        self.assertEqual(html.count("<h1"), 1)

    def test_drip_tape_catalog_title_and_heading(self):
        Category.objects.create(name="نوار تیپ", slug="drip-tape", tagline="نوار تیپ آبیاری")
        response = self.client.get(reverse("catalog") + "?category=drip-tape")
        self.assertContains(response, "خرید نوار تیپ | کشت‌دریپ")
        self.assertContains(response, "<h1")
        self.assertContains(response, "خرید نوار تیپ")
        self.assertContains(response, "index, follow")

    def test_shop_and_search_robots(self):
        shop = self.client.get(reverse("catalog"))
        self.assertContains(shop, "خرید لوازم آبیاری | کشت‌دریپ")
        search = self.client.get(reverse("catalog") + "?q=نوار")
        self.assertContains(search, "noindex, follow")
        cart = self.client.get(reverse("cart"))
        self.assertContains(cart, "noindex, follow")

    def test_robots_and_sitemap(self):
        category = Category.objects.create(name="نوار تیپ", slug="drip-tape", tagline="نوار")
        brand = Brand.objects.create(name="کشت‌دریپ", slug="seo-brand")
        Product.objects.create(
            category=category,
            brand=brand,
            name="نوار تیپ نمونه",
            slug="seo-tape",
            blurb="نوار تیپ",
            description="توضیح",
            price=1000,
            unit="رول",
        )
        robots = self.client.get(reverse("robots"))
        self.assertContains(robots, "Sitemap:")
        self.assertContains(robots, "Disallow: /basket/")
        sitemap = self.client.get(reverse("sitemap"))
        self.assertContains(sitemap, "category=drip-tape")
        self.assertContains(sitemap, "/product/seo-tape/")
        self.assertContains(sitemap, "/blog/how-to-choose-drip-tape/")

    def test_product_schema_uses_rial(self):
        category = Category.objects.create(name="نوار تیپ", slug="seo-tape-cat", tagline="نوار")
        brand = Brand.objects.create(name="کشت‌دریپ", slug="seo-brand-2")
        product = Product.objects.create(
            category=category,
            brand=brand,
            name="نوار تیپ تست",
            slug="seo-tape-product",
            blurb="نوار تیپ برای آبیاری",
            description="توضیح",
            price=250000,
            unit="رول",
        )
        response = self.client.get(product.get_absolute_url())
        self.assertContains(response, "خرید نوار تیپ تست | کشت‌دریپ")
        self.assertContains(response, '"priceCurrency": "IRR"')
        self.assertContains(response, '"price": 2500000')


def price_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "prices.xlsx"
    buffer.size = buffer.getbuffer().nbytes
    return buffer


class PriceImportTests(TestCase):
    def setUp(self):
        category = Category.objects.create(name="آبیاری", slug="import-irrigation", tagline="آب")
        brand = Brand.objects.create(name="برند", slug="import-brand")
        self.product = Product.objects.create(
            category=category,
            brand=brand,
            name="لوله قطره‌ای",
            slug="KSH-P-00000044",
            blurb="لوله",
            description="توضیح",
            price=1000,
            unit="متر",
        )
        self.variant = ProductVariant.objects.create(
            product=self.product,
            slug="KSH-V-00000044",
            working_pressure="4 بار",
            diameter="16",
            price=100000,
        )
        self.plain = Product.objects.create(
            category=category,
            brand=brand,
            name="کوزه",
            slug="olla",
            blurb="کوزه",
            description="توضیح",
            price=640000,
            unit="جفت",
        )

    def test_import_updates_product_and_variant_prices(self):
        upload = price_workbook(
            [
                ["کد محصول", "نام", "فشار کاری", "قطر", "قیمت"],
                ["olla", "کوزه", "", "", 700000],
                ["KSH-P-00000044", "لوله قطره‌ای", "4 بار", 16, "۱٬۲۵۰٬۰۰۰"],
            ]
        )
        result = import_price_workbook(upload)
        self.assertEqual(result.updated_products, 1)
        self.assertEqual(result.updated_variants, 1)
        self.assertEqual(result.errors, [])
        self.plain.refresh_from_db()
        self.variant.refresh_from_db()
        self.assertEqual(self.plain.price, 700000)
        self.assertEqual(self.variant.price, 1250000)

    def test_unknown_code_and_missing_variant_are_reported(self):
        upload = price_workbook(
            [
                ["slug", "price", "working_pressure", "diameter"],
                ["missing", 10, "", ""],
                ["olla", 640000, "", ""],
                ["KSH-P-00000044", 20, "9 بار", ""],
            ]
        )
        result = import_price_workbook(upload)
        self.assertEqual(result.unchanged, 1)
        self.assertEqual(len(result.errors), 2)
        self.plain.refresh_from_db()
        self.variant.refresh_from_db()
        self.assertEqual(self.plain.price, 640000)
        self.assertEqual(self.variant.price, 100000)

    def test_import_updates_a_variant_by_its_slug(self):
        upload = price_workbook(
            [
                ["کد محصول", "قیمت"],
                ["KSH-V-00000044", 900000],
            ]
        )
        result = import_price_workbook(upload)
        self.assertEqual(result.updated_variants, 1)
        self.assertEqual(result.errors, [])
        self.variant.refresh_from_db()
        self.assertEqual(self.variant.price, 900000)

    def test_variant_row_without_slug_is_rejected(self):
        upload = price_workbook(
            [
                ["نام", "فشار کاری", "قطر", "قیمت"],
                ["لوله قطره‌ای", "4 بار", "16", 20],
            ]
        )
        result = import_price_workbook(upload)
        self.assertEqual(result.updated_variants, 0)
        self.assertEqual(len(result.errors), 1)
        self.variant.refresh_from_db()
        self.assertEqual(self.variant.price, 100000)

    def test_rejects_a_file_that_is_not_xlsx(self):
        upload = BytesIO(b"not excel")
        upload.name = "prices.csv"
        upload.size = 9
        with self.assertRaises(PriceImportError):
            import_price_workbook(upload)

    def test_admin_import_page_updates_prices(self):
        user = get_user_model().objects.create_superuser("admin", "admin@example.com", "secret-pass")
        self.client.force_login(user)
        page = self.client.get(reverse("admin:products_product_import_prices"))
        self.assertContains(page, "ورود قیمت از اکسل")
        self.assertContains(page, "دانلود قیمت‌های فعلی")

        changelist = self.client.get(reverse("admin:products_product_changelist"))
        self.assertContains(changelist, "ورود قیمت از اکسل")

        upload = price_workbook(
            [
                ["کد محصول", "قیمت"],
                ["olla", 810000],
            ]
        )
        response = self.client.post(reverse("admin:products_product_import_prices"), {"file": upload})
        self.assertContains(response, "1 قیمت محصول")
        self.plain.refresh_from_db()
        self.assertEqual(self.plain.price, 810000)

        exported = self.client.get(reverse("admin:products_product_export_prices"))
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(BytesIO(exported.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "کد محصول")
        exported_prices = {row[0]: row[4] for row in rows[1:]}
        self.assertEqual(exported_prices["olla"], 810000)
        self.assertEqual(exported_prices["KSH-V-00000044"], 100000)

    def test_anonymous_user_cannot_import_prices(self):
        response = self.client.get(reverse("admin:products_product_import_prices"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response.url)
