from collections import defaultdict
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .models import Product, ProductVariant

HEADERS = ["کد محصول", "نام", "فشار کاری", "قطر", "قیمت"]
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_PRICE = 2_147_483_647

COLUMN_ALIASES = {
    "slug": "slug",
    "اسلاگ": "slug",
    "نامک": "slug",
    "کد": "slug",
    "کد محصول": "slug",
    "id": "id",
    "شناسه": "id",
    "name": "name",
    "نام": "name",
    "نام محصول": "name",
    "price": "price",
    "قیمت": "price",
    "قیمت تومان": "price",
    "working pressure": "working_pressure",
    "فشار کاری": "working_pressure",
    "فشار": "working_pressure",
    "diameter": "diameter",
    "قطر": "diameter",
}

DIGIT_TRANSLATION = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class PriceImportError(Exception):
    pass


class PriceImportResult:
    def __init__(self):
        self.updated_products = 0
        self.updated_variants = 0
        self.unchanged = 0
        self.errors = []

    @property
    def updated(self):
        return self.updated_products + self.updated_variants

    def summary(self):
        parts = []
        if self.updated_products:
            parts.append(f"{self.updated_products} قیمت محصول")
        if self.updated_variants:
            parts.append(f"{self.updated_variants} قیمت تنوع")
        if parts:
            text = "به‌روزرسانی شد: " + " و ".join(parts) + "."
        else:
            text = "هیچ قیمتی تغییر نکرد."
        if self.unchanged:
            text += f" {self.unchanged} ردیف همان قیمت قبلی را داشت."
        if self.errors:
            text += f" {len(self.errors)} ردیف خطا داشت."
        return text


def build_price_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "قیمت‌ها"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    products = Product.objects.prefetch_related("variants").order_by("name", "id")
    for product in products:
        variants = list(product.variants.all())
        if variants:
            for variant in variants:
                sheet.append(
                    [
                        variant.slug,
                        product.name,
                        variant.working_pressure,
                        variant.diameter,
                        variant.price,
                    ]
                )
        else:
            sheet.append([product.slug, product.name, "", "", product.price])

    sheet.freeze_panes = "A2"
    for index, width in enumerate((22, 36, 18, 18, 16), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if sheet.max_row >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    return workbook


def workbook_bytes():
    buffer = BytesIO()
    build_price_workbook().save(buffer)
    return buffer.getvalue()


def import_price_workbook(upload):
    filename = (getattr(upload, "name", "") or "").lower()
    if filename and not filename.endswith(".xlsx"):
        raise PriceImportError("فقط فایل اکسل با پسوند xlsx پذیرفته می‌شود.")
    size = getattr(upload, "size", None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        raise PriceImportError("حجم فایل بیشتر از ۵ مگابایت است.")

    try:
        workbook = load_workbook(upload, data_only=True, read_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise PriceImportError("فایل اکسل معتبر نیست.") from exc

    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index, columns = _header_columns(rows)
    data_rows = rows[header_index + 1 :]
    if len(data_rows) > MAX_ROWS:
        raise PriceImportError(f"فایل بیشتر از {MAX_ROWS} ردیف قیمت دارد.")

    products = list(Product.objects.prefetch_related("variants"))
    by_slug = {product.slug: product for product in products}
    by_variant_slug = {
        variant.slug: variant
        for variant in ProductVariant.objects.select_related("product")
    }
    by_id = {product.pk: product for product in products}
    by_name = defaultdict(list)
    for product in products:
        by_name[product.name].append(product)

    result = PriceImportResult()
    with transaction.atomic():
        for offset, row in enumerate(data_rows, start=header_index + 2):
            if _row_is_empty(row):
                continue
            try:
                _apply_row(row, columns, by_slug, by_id, by_name, by_variant_slug, result)
            except PriceImportError as exc:
                result.errors.append(f"ردیف {offset}: {exc}")
    return result


def _header_columns(rows):
    for index, row in enumerate(rows):
        if _row_is_empty(row):
            continue
        columns = {}
        for column_index, value in enumerate(row):
            key = COLUMN_ALIASES.get(_normalize_header(value))
            if key and key not in columns:
                columns[key] = column_index
        if "price" in columns and ({"slug", "id", "name"} & set(columns)):
            return index, columns
        raise PriceImportError(
            "ردیف عنوان باید ستون قیمت و حداقل یکی از ستون‌های کد محصول، شناسه یا نام را داشته باشد."
        )
    raise PriceImportError("فایل خالی است.")


def _apply_row(row, columns, by_slug, by_id, by_name, by_variant_slug, result):
    price = _parse_price(_cell(row, columns.get("price")))
    slug = _cell_text(_cell(row, columns.get("slug")))
    variant = by_variant_slug.get(slug) if slug else None
    if variant is not None:
        if variant.price == price:
            result.unchanged += 1
            return
        variant.price = price
        variant.save(update_fields=["price"])
        result.updated_variants += 1
        return

    product = _find_product(row, columns, by_slug, by_id, by_name)
    pressure = _cell_text(_cell(row, columns.get("working_pressure")))
    diameter = _cell_text(_cell(row, columns.get("diameter")))

    if pressure or diameter:
        if not slug:
            raise PriceImportError("کد را برای این تنوع وارد کنید.")
        variant = _find_variant(product, pressure, diameter)
        if variant.price == price:
            result.unchanged += 1
            return
        variant.price = price
        variant.save(update_fields=["price"])
        result.updated_variants += 1
        return

    if product.price == price:
        result.unchanged += 1
        return
    product.price = price
    product.save(update_fields=["price"])
    result.updated_products += 1


def _find_product(row, columns, by_slug, by_id, by_name):
    slug = _cell_text(_cell(row, columns.get("slug")))
    if slug:
        product = by_slug.get(slug)
        if product is None:
            raise PriceImportError(f"محصولی با کد «{slug}» پیدا نشد.")
        return product

    raw_id = _cell(row, columns.get("id"))
    if raw_id not in (None, ""):
        try:
            product_id = int(_cell_text(raw_id))
        except ValueError as exc:
            raise PriceImportError("شناسه محصول باید عدد باشد.") from exc
        product = by_id.get(product_id)
        if product is None:
            raise PriceImportError(f"محصولی با شناسه {product_id} پیدا نشد.")
        return product

    name = _cell_text(_cell(row, columns.get("name")))
    if not name:
        raise PriceImportError("کد محصول، شناسه یا نام را وارد کنید.")
    matches = by_name.get(name, [])
    if not matches:
        raise PriceImportError(f"محصولی با نام «{name}» پیدا نشد.")
    if len(matches) > 1:
        raise PriceImportError(f"چند محصول با نام «{name}» وجود دارد. کد محصول را وارد کنید.")
    return matches[0]


def _find_variant(product, pressure, diameter):
    for variant in product.variants.all():
        if variant.working_pressure == pressure and variant.diameter == diameter:
            return variant
    details = []
    if pressure:
        details.append(f"فشار کاری «{pressure}»")
    if diameter:
        details.append(f"قطر «{diameter}»")
    spec = " و ".join(details)
    raise PriceImportError(f"برای «{product.name}» تنوعی با {spec} پیدا نشد.")


def _parse_price(value):
    if isinstance(value, bool) or value is None or (isinstance(value, str) and not value.strip()):
        raise PriceImportError("قیمت خالی است.")
    if isinstance(value, float):
        if not value.is_integer():
            raise PriceImportError("قیمت باید عدد صحیح تومان باشد.")
        number = int(value)
    elif isinstance(value, int):
        number = value
    else:
        text = str(value).translate(DIGIT_TRANSLATION)
        for token in (",", "،", "٬", " ", "تومان"):
            text = text.replace(token, "")
        text = text.strip()
        if not text.isdigit():
            raise PriceImportError(f"قیمت «{value}» قابل خواندن نیست.")
        number = int(text)
    if number < 0:
        raise PriceImportError("قیمت نمی‌تواند منفی باشد.")
    if number > MAX_PRICE:
        raise PriceImportError("قیمت از حد مجاز بزرگ‌تر است.")
    return number


def _cell(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _row_is_empty(row):
    return row is None or all(value is None or str(value).strip() == "" for value in row)


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).replace("\u200c", " ").strip().lower()
    text = text.replace("_", " ").replace("-", " ")
    return " ".join(text.split())
